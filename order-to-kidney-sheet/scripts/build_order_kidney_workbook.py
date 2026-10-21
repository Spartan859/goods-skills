#!/usr/bin/env python3
"""Build the order allocation and kidney summary xlsm from normalized JSON."""

from __future__ import annotations

import argparse
import json
import math
import os
import sys
from copy import copy
from pathlib import Path
from typing import Any

from openpyxl import load_workbook
from openpyxl.styles import Alignment
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table, TableColumn

KIDNEY_HEADERS = ["CN", "应付", "点数"]
TOP_TABLE_NAME = "海边心_排表32"
BOTTOM_TABLE_NAME = "表9"
DEFAULT_ALLOCATION_COUNT = 10
DEFAULT_TEMPLATE = Path(__file__).resolve().parents[1] / "assets" / "order-kidney-template.xlsm"

SPEC_KEYS = {"source_mode", "title", "allocation_count", "products", "kidney_rows"}
PRODUCT_KEYS = {
    "name",
    "allocations",
    "unit_price",
    "exchange_rate",
    "adjustment",
    "price",
}
KIDNEY_ROW_KEYS = {"cn", "amount", "points"}


class WorkbookBuildError(ValueError):
    """Raised when normalized input cannot produce a safe workbook."""


def _plain_mapping(value: Any, label: str) -> dict[str, Any]:
    if not isinstance(value, dict):
        raise WorkbookBuildError(f"{label} 必须是对象")
    return value


def _reject_unknown(mapping: dict[str, Any], allowed: set[str], label: str) -> None:
    unknown = sorted(set(mapping) - allowed)
    if unknown:
        raise WorkbookBuildError(f"{label} 包含未知字段: {', '.join(unknown)}")


def _finite_number(value: Any, label: str) -> float:
    if isinstance(value, bool) or not isinstance(value, (int, float)):
        raise WorkbookBuildError(f"{label} 必须是数字")
    number = float(value)
    if not math.isfinite(number):
        raise WorkbookBuildError(f"{label} 必须是有限数字")
    return number


def _nonempty_text(value: Any, label: str) -> str:
    if not isinstance(value, str) or not value.strip():
        raise WorkbookBuildError(f"{label} 必须是非空字符串")
    return value.strip()


def _optional_number(product: dict[str, Any], key: str, label: str) -> float | None:
    value = product.get(key)
    if value is None:
        return None
    return _finite_number(value, label)


def _allocation_count(spec: dict[str, Any]) -> int:
    raw_count = spec.get("allocation_count", DEFAULT_ALLOCATION_COUNT)
    if isinstance(raw_count, bool) or not isinstance(raw_count, int) or raw_count < 1:
        raise WorkbookBuildError("allocation_count 必须是大于 0 的整数")
    return raw_count


def _allocation_headers(allocation_count: int) -> list[str]:
    return [f"第{i}配" for i in range(1, allocation_count + 1)]


def _top_headers(allocation_count: int) -> list[str]:
    return ["列1", *_allocation_headers(allocation_count), "调价", "价格", "汇率", "单价"]


def _validate_spec(raw: Any) -> dict[str, Any]:
    spec = _plain_mapping(raw, "输入")
    _reject_unknown(spec, SPEC_KEYS, "输入")

    mode = spec.get("source_mode")
    if mode not in {"arrangement", "kidney"}:
        raise WorkbookBuildError("source_mode 必须是 arrangement 或 kidney")

    title = spec.get("title", "")
    if title is None:
        title = ""
    if not isinstance(title, str):
        raise WorkbookBuildError("title 必须是字符串")
    title = title.strip()
    allocation_count = _allocation_count(spec)

    products_raw = spec.get("products", [])
    kidney_raw = spec.get("kidney_rows", [])
    if not isinstance(products_raw, list):
        raise WorkbookBuildError("products 必须是数组")
    if not isinstance(kidney_raw, list):
        raise WorkbookBuildError("kidney_rows 必须是数组")

    products: list[dict[str, Any]] = []
    for index, raw_product in enumerate(products_raw, start=1):
        product = _plain_mapping(raw_product, f"products[{index}]")
        _reject_unknown(product, PRODUCT_KEYS, f"products[{index}]")
        name = _nonempty_text(product.get("name"), f"products[{index}].name")

        allocations = product.get("allocations", [])
        if not isinstance(allocations, list):
            raise WorkbookBuildError(f"products[{index}].allocations 必须是数组")
        if len(allocations) > allocation_count:
            raise WorkbookBuildError(
                f"products[{index}].allocations 超过 allocation_count={allocation_count}"
            )
        normalized_allocations: list[str | None] = []
        for slot_index, value in enumerate(allocations, start=1):
            if value is None or value == "":
                normalized_allocations.append(None)
            elif isinstance(value, str):
                normalized_allocations.append(value.strip() or None)
            else:
                raise WorkbookBuildError(
                    f"products[{index}].allocations[{slot_index}] 必须是用户名或 null"
                )
        normalized_allocations.extend(
            [None] * (allocation_count - len(normalized_allocations))
        )

        price = _optional_number(product, "price", f"products[{index}].price")
        unit_price = _optional_number(
            product, "unit_price", f"products[{index}].unit_price"
        )
        exchange_rate = _optional_number(
            product, "exchange_rate", f"products[{index}].exchange_rate"
        )
        adjustment_value = product.get("adjustment")
        adjustment = _optional_number(
            product, "adjustment", f"products[{index}].adjustment"
        )
        if adjustment is None:
            adjustment = 0.0

        if any(normalized_allocations):
            if price is not None:
                if (
                    unit_price is not None
                    or exchange_rate is not None
                    or adjustment_value is not None
                ):
                    raise WorkbookBuildError(
                        f"products[{index}] 不能同时提供 price 和排表价格计算参数"
                    )
            elif unit_price is None or exchange_rate is None:
                raise WorkbookBuildError(
                    f"products[{index}] 需要 price，或同时提供 unit_price 与 exchange_rate"
                )

        products.append(
            {
                "name": name,
                "allocations": normalized_allocations,
                "price": price,
                "unit_price": unit_price,
                "exchange_rate": exchange_rate,
                "adjustment": adjustment,
                "row_index": index + 2,
            }
        )

    kidney_rows: list[dict[str, Any]] = []
    for index, raw_row in enumerate(kidney_raw, start=1):
        row = _plain_mapping(raw_row, f"kidney_rows[{index}]")
        _reject_unknown(row, KIDNEY_ROW_KEYS, f"kidney_rows[{index}]")
        cn = _nonempty_text(row.get("cn"), f"kidney_rows[{index}].cn")
        amount = _finite_number(row.get("amount"), f"kidney_rows[{index}].amount")
        points_value = row.get("points")
        points = None
        if points_value is not None:
            points_number = _finite_number(
                points_value, f"kidney_rows[{index}].points"
            )
            if points_number < 0 or not points_number.is_integer():
                raise WorkbookBuildError(
                    f"kidney_rows[{index}].points 必须是非负整数"
                )
            points = int(points_number)
        kidney_rows.append({"cn": cn, "amount": amount, "points": points})

    if mode == "arrangement":
        if not products:
            raise WorkbookBuildError("arrangement 模式要求至少一个商品")
        if kidney_rows:
            raise WorkbookBuildError("arrangement 模式不应提供 kidney_rows")
        if not any(any(product["allocations"]) for product in products):
            raise WorkbookBuildError("arrangement 模式要求至少一个非空配位")
    else:
        if products:
            raise WorkbookBuildError("kidney 模式不应提供 products")
        if not kidney_rows:
            raise WorkbookBuildError("kidney 模式要求至少一行 kidney_rows")

    return {
        "source_mode": mode,
        "title": title,
        "allocation_count": allocation_count,
        "products": products,
        "kidney_rows": kidney_rows,
    }


def _product_price(product: dict[str, Any]) -> float:
    if product["price"] is not None:
        return float(product["price"])
    return (
        float(product["unit_price"])
        * float(product["exchange_rate"])
        + float(product["adjustment"])
    )


def _arrangement_summary(
    products: list[dict[str, Any]], allocation_count: int
) -> list[dict[str, Any]]:
    order: list[str] = []
    totals: dict[str, dict[str, Any]] = {}
    for slot_index in range(allocation_count):
        for product in products:
            cn = product["allocations"][slot_index]
            if not cn:
                continue
            if cn not in totals:
                order.append(cn)
                totals[cn] = {"cn": cn, "amount": 0.0, "points": 0}
            totals[cn]["amount"] += _product_price(product)
            totals[cn]["points"] += 1
    return [totals[cn] for cn in order]


def _kidney_summary(rows: list[dict[str, Any]]) -> list[dict[str, Any]]:
    order: list[str] = []
    totals: dict[str, dict[str, Any]] = {}
    for row in rows:
        cn = row["cn"]
        if cn not in totals:
            order.append(cn)
            totals[cn] = {"cn": cn, "amount": 0.0, "points": None}
        totals[cn]["amount"] += row["amount"]
        if row["points"] is not None:
            current = totals[cn]["points"] or 0
            totals[cn]["points"] = current + row["points"]
    return [totals[cn] for cn in order]


def _copy_cell_style(source, target) -> None:
    if source.has_style:
        target._style = copy(source._style)
    if source.number_format:
        target.number_format = source.number_format
    if source.alignment:
        target.alignment = copy(source.alignment)
    if source.protection:
        target.protection = copy(source.protection)


def _top_template_column(column: int, allocation_count: int) -> int:
    if column == 1:
        return 1
    if column <= allocation_count + 1:
        return 2
    return 11 + (column - (allocation_count + 1))


def _prepare_layout(
    ws, product_count: int, kidney_count: int, allocation_count: int
) -> tuple[int, int, int, int]:
    top_end = 3 + max(product_count, 1) - 1
    lower_title_row = max(11, top_end + 3)
    bottom_header_row = lower_title_row + 1
    bottom_end_row = bottom_header_row + max(kidney_count, 1)
    last_column = allocation_count + 5

    for merged_range in list(ws.merged_cells.ranges):
        if (
            merged_range.min_row == 1
            and merged_range.min_col == 1
            or str(merged_range).startswith("A11:C11")
        ):
            ws.unmerge_cells(str(merged_range))

    max_clear_row = max(28, bottom_end_row + 2)
    max_clear_column = max(21, last_column)
    for row in ws.iter_rows(
        min_row=1,
        max_row=max_clear_row,
        min_col=1,
        max_col=max_clear_column,
    ):
        for cell in row:
            cell.value = None

    for target_row in range(2, top_end + 1):
        source_row = 2 if target_row == 2 else 3
        for target_column in range(1, last_column + 1):
            source_column = _top_template_column(target_column, allocation_count)
            _copy_cell_style(
                ws.cell(source_row, source_column),
                ws.cell(target_row, target_column),
            )
        if target_row not in (2, 3):
            ws.row_dimensions[target_row].height = ws.row_dimensions[3].height

    if lower_title_row != 11:
        for column in range(1, 4):
            _copy_cell_style(ws.cell(11, column), ws.cell(lower_title_row, column))
    for target_row, source_row in (
        (bottom_header_row, 12),
        *[(r, 13) for r in range(bottom_header_row + 1, bottom_end_row + 1)],
    ):
        if target_row != source_row:
            for column in range(1, 4):
                _copy_cell_style(ws.cell(source_row, column), ws.cell(target_row, column))
            ws.row_dimensions[target_row].height = ws.row_dimensions[source_row].height

    ws.merge_cells(f"A1:{get_column_letter(last_column)}1")
    ws.merge_cells(f"A{lower_title_row}:C{lower_title_row}")
    ws["A1"].alignment = Alignment(horizontal="left", vertical="center")
    ws[f"A{lower_title_row}"].alignment = Alignment(
        horizontal="left", vertical="center"
    )
    return top_end, lower_title_row, bottom_header_row, bottom_end_row


def _set_table_ref(table: Table, ref: str) -> None:
    table.ref = ref
    if table.autoFilter is not None:
        table.autoFilter.ref = ref


def _clear_table_formulas(ws) -> None:
    for table in ws.tables.values():
        for column in table.tableColumns:
            column.calculatedColumnFormula = None


def _set_table_columns(table: Table, headers: list[str]) -> None:
    table.tableColumns = [
        TableColumn(id=index, name=header)
        for index, header in enumerate(headers, start=1)
    ]


def build_workbook(
    raw_spec: Any, output_path: str | Path, template_path: str | Path | None = None
) -> Path:
    spec = _validate_spec(raw_spec)
    output = Path(output_path).expanduser().resolve()
    if output.suffix.lower() != ".xlsm":
        raise WorkbookBuildError("输出文件必须使用 .xlsm 扩展名")
    template = (
        Path(template_path).expanduser().resolve()
        if template_path is not None
        else DEFAULT_TEMPLATE.resolve()
    )
    if not template.is_file():
        raise WorkbookBuildError(f"模板不存在: {template}")

    wb = load_workbook(template, keep_vba=True)
    if "Sheet1" not in wb.sheetnames:
        raise WorkbookBuildError("模板必须包含 Sheet1")
    ws = wb["Sheet1"]
    ws._images.clear()

    products = spec["products"]
    allocation_count = spec["allocation_count"]
    kidney_rows = (
        _arrangement_summary(products, allocation_count)
        if spec["source_mode"] == "arrangement"
        else _kidney_summary(spec["kidney_rows"])
    )
    top_end, lower_title_row, bottom_header_row, bottom_end_row = _prepare_layout(
        ws, len(products), len(kidney_rows), allocation_count
    )

    title = spec["title"]
    ws["A1"] = f"{title} 排表".strip()
    ws[f"A{lower_title_row}"] = f"{title} 肾表".strip()

    top_headers = _top_headers(allocation_count)
    top_last_column = len(top_headers)
    adjustment_column = allocation_count + 2
    price_column = adjustment_column + 1
    exchange_column = price_column + 1
    unit_column = exchange_column + 1
    for column, header in enumerate(top_headers, start=1):
        ws.cell(2, column).value = header
    for column, header in enumerate(KIDNEY_HEADERS, start=1):
        ws.cell(bottom_header_row, column).value = header

    if spec["source_mode"] == "arrangement":
        for offset, product in enumerate(products):
            row = 3 + offset
            ws.cell(row, 1).value = product["name"]
            for slot_offset, cn in enumerate(product["allocations"], start=2):
                ws.cell(row, slot_offset).value = cn
            ws.cell(row, adjustment_column).value = product["adjustment"]
            if product["price"] is not None:
                ws.cell(row, price_column).value = product["price"]
            else:
                adjustment_cell = f"{get_column_letter(adjustment_column)}{row}"
                exchange_cell = f"{get_column_letter(exchange_column)}{row}"
                unit_cell = f"{get_column_letter(unit_column)}{row}"
                ws.cell(row, price_column).value = (
                    f"={unit_cell}*{exchange_cell}+{adjustment_cell}"
                )
            ws.cell(row, exchange_column).value = product["exchange_rate"]
            ws.cell(row, unit_column).value = product["unit_price"]
            ws.cell(row, price_column).number_format = "0.00"
    else:
        for column in range(1, top_last_column + 1):
            ws.cell(3, column).value = None

    for offset, item in enumerate(kidney_rows):
        row = bottom_header_row + 1 + offset
        ws.cell(row, 1).value = item["cn"]
        ws.cell(row, 2).value = item["amount"]
        ws.cell(row, 2).number_format = "0.00"
        ws.cell(row, 3).value = item["points"]

    top_last_letter = get_column_letter(top_last_column)
    top_table = ws.tables[TOP_TABLE_NAME]
    _set_table_ref(top_table, f"A2:{top_last_letter}{top_end}")
    _set_table_columns(top_table, top_headers)
    _set_table_ref(
        ws.tables[BOTTOM_TABLE_NAME],
        f"A{bottom_header_row}:C{bottom_end_row}",
    )
    _clear_table_formulas(ws)

    wb.calculation.calcMode = "auto"
    wb.calculation.fullCalcOnLoad = True
    wb.calculation.forceFullCalc = True
    output.parent.mkdir(parents=True, exist_ok=True)
    temp_output = output.with_name(f".{output.stem}.tmp.xlsm")
    try:
        wb.save(temp_output)
        os.replace(temp_output, output)
    finally:
        temp_output.unlink(missing_ok=True)
    return output


def main() -> int:
    parser = argparse.ArgumentParser(
        description="从规范化 JSON 生成排表/肾表 xlsm"
    )
    parser.add_argument("--input", required=True, help="规范化 JSON 文件")
    parser.add_argument("--output", required=True, help="输出 xlsm 文件")
    parser.add_argument("--template", help="可选的自定义 xlsm 模板")
    args = parser.parse_args()

    try:
        with Path(args.input).expanduser().open("r", encoding="utf-8") as handle:
            raw_spec = json.load(handle)
        build_workbook(raw_spec, args.output, args.template)
    except (OSError, json.JSONDecodeError, WorkbookBuildError) as error:
        print(f"生成失败: {error}", file=sys.stderr)
        return 2
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
