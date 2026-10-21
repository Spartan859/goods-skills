import sys
import tempfile
import unittest
import zipfile
from pathlib import Path

from openpyxl import load_workbook

SKILL_ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(SKILL_ROOT / "scripts"))

from build_order_kidney_workbook import build_workbook


class ArrangementModeTests(unittest.TestCase):
    def test_builds_arrangement_prices_and_kidney_summary(self):
        spec = {
            "source_mode": "arrangement",
            "title": "测试团",
            "products": [
                {
                    "name": "梨",
                    "allocations": ["阿甲", None, None, None, None, None, None, None, None, None],
                    "unit_price": 400,
                    "exchange_rate": 0.0455,
                    "adjustment": 6,
                },
                {
                    "name": "善",
                    "allocations": ["阿甲", "阿庚", None, None, None, None, None, None, None, None],
                    "unit_price": 400,
                    "exchange_rate": 0.0455,
                    "adjustment": 12,
                },
            ],
            "kidney_rows": [],
        }
        with tempfile.TemporaryDirectory() as temp_dir:
            out = Path(temp_dir) / "arrangement.xlsm"
            build_workbook(spec, out)
            wb = load_workbook(out, keep_vba=True, data_only=False)
            ws = wb["Sheet1"]

            self.assertEqual(ws["A1"].value, "测试团 排表")
            self.assertEqual(ws["A11"].value, "测试团 肾表")
            self.assertEqual([ws["A3"].value, ws["A4"].value], ["梨", "善"])
            self.assertEqual([ws["B3"].value, ws["B4"].value, ws["C4"].value], ["阿甲", "阿甲", "阿庚"])
            self.assertEqual([ws["L3"].value, ws["N3"].value, ws["O3"].value], [6, 0.0455, 400])
            self.assertEqual(ws["M3"].value, "=O3*N3+L3")
            self.assertEqual(ws["M4"].value, "=O4*N4+L4")

            self.assertEqual([ws["A12"].value, ws["B12"].value, ws["C12"].value], ["CN", "应付", "点数"])
            self.assertEqual(ws["A13"].value, "阿甲")
            self.assertAlmostEqual(ws["B13"].value, 54.4)
            self.assertEqual(ws["C13"].value, 2)
            self.assertEqual(ws["A14"].value, "阿庚")
            self.assertAlmostEqual(ws["B14"].value, 30.2)
            self.assertEqual(ws["C14"].value, 1)

            self.assertIn("A2:O4", [table.ref for table in ws.tables.values()])
            self.assertIn("A12:C14", [table.ref for table in ws.tables.values()])
            self.assertIsNotNone(wb.vba_archive)
            self.assertEqual(len(ws._images), 0)


class KidneyModeTests(unittest.TestCase):
    def test_builds_three_column_kidney_sheet_without_arrangement_data(self):
        spec = {
            "source_mode": "kidney",
            "title": "示例团",
            "products": [],
            "kidney_rows": [
                {"cn": "甲", "amount": 41.57, "points": None},
                {"cn": "乙", "amount": 10.71, "points": None},
                {"cn": "甲", "amount": 2.43, "points": 2},
            ],
        }
        with tempfile.TemporaryDirectory() as temp_dir:
            out = Path(temp_dir) / "kidney.xlsm"
            build_workbook(spec, out)
            wb = load_workbook(out, keep_vba=True, data_only=False)
            ws = wb["Sheet1"]

            self.assertEqual(ws["A1"].value, "示例团 排表")
            self.assertEqual(ws["A3"].value, None)
            self.assertTrue(all(ws.cell(3, column).value is None for column in range(1, 16)))
            self.assertEqual([ws["A12"].value, ws["B12"].value, ws["C12"].value], ["CN", "应付", "点数"])
            self.assertEqual(ws["A13"].value, "甲")
            self.assertAlmostEqual(ws["B13"].value, 44.0)
            self.assertEqual(ws["C13"].value, 2)
            self.assertEqual(ws["A14"].value, "乙")
            self.assertAlmostEqual(ws["B14"].value, 10.71)
            self.assertIsNone(ws["C14"].value)
            self.assertIn("A12:C14", [table.ref for table in ws.tables.values()])
            self.assertIsNotNone(wb.vba_archive)
            self.assertEqual(len(ws._images), 0)


class LayoutAndValidationTests(unittest.TestCase):
    def test_moves_kidney_section_below_more_than_six_products(self):
        products = []
        for index in range(8):
            products.append(
                {
                    "name": f"商品{index + 1}",
                    "allocations": ["用户8", None, None, None, None, None, None, None, None, None],
                    "price": 10,
                }
            )
        spec = {
            "source_mode": "arrangement",
            "title": "动态布局",
            "products": products,
            "kidney_rows": [],
        }
        with tempfile.TemporaryDirectory() as temp_dir:
            out = Path(temp_dir) / "eight-products.xlsm"
            build_workbook(spec, out)
            wb = load_workbook(out, keep_vba=True, data_only=False)
            ws = wb["Sheet1"]

            self.assertEqual(ws["A10"].value, "商品8")
            self.assertEqual(ws["A13"].value, "动态布局 肾表")
            self.assertEqual([ws["A14"].value, ws["B14"].value, ws["C14"].value], ["CN", "应付", "点数"])
            self.assertEqual(ws["A15"].value, "用户8")
            self.assertAlmostEqual(ws["B15"].value, 80.0)
            self.assertEqual(ws["C15"].value, 8)
            self.assertIn("A2:O10", [table.ref for table in ws.tables.values()])
            self.assertIn("A14:C15", [table.ref for table in ws.tables.values()])

    def test_rejects_more_than_ten_allocations(self):
        spec = {
            "source_mode": "arrangement",
            "title": "错误",
            "products": [
                {
                    "name": "梨",
                    "allocations": ["A"] * 11,
                    "price": 1,
                }
            ],
            "kidney_rows": [],
        }
        with tempfile.TemporaryDirectory() as temp_dir:
            out = Path(temp_dir) / "invalid.xlsm"
            with self.assertRaisesRegex(ValueError, "最多包含 10 个配位"):
                build_workbook(spec, out)
            self.assertFalse(out.exists())

    def test_rejects_arrangement_without_price_information(self):
        spec = {
            "source_mode": "arrangement",
            "title": "缺价",
            "products": [
                {
                    "name": "梨",
                    "allocations": ["A", None, None, None, None, None, None, None, None, None],
                }
            ],
            "kidney_rows": [],
        }
        with tempfile.TemporaryDirectory() as temp_dir:
            out = Path(temp_dir) / "missing-price.xlsm"
            with self.assertRaisesRegex(ValueError, "需要 price"):
                build_workbook(spec, out)
            self.assertFalse(out.exists())

    def test_rejects_fixed_price_with_adjustment(self):
        spec = {
            "source_mode": "arrangement",
            "title": "冲突",
            "products": [
                {
                    "name": "梨",
                    "allocations": ["A", None, None, None, None, None, None, None, None, None],
                    "price": 12.5,
                    "adjustment": 1,
                }
            ],
            "kidney_rows": [],
        }
        with tempfile.TemporaryDirectory() as temp_dir:
            out = Path(temp_dir) / "conflicting-price.xlsm"
            with self.assertRaisesRegex(ValueError, "不能同时提供"):
                build_workbook(spec, out)
            self.assertFalse(out.exists())

    def test_rejects_unknown_fields_and_invalid_amounts(self):
        cases = [
            (
                "未知顶层字段",
                {
                    "source_mode": "kidney",
                    "products": [],
                    "kidney_rows": [{"cn": "甲", "amount": 1}],
                    "extra": True,
                },
            ),
            (
                "未知商品字段",
                {
                    "source_mode": "arrangement",
                    "products": [
                        {
                            "name": "梨",
                            "allocations": ["A"],
                            "price": 1,
                            "extra": True,
                        }
                    ],
                    "kidney_rows": [],
                },
            ),
            (
                "非数字金额",
                {
                    "source_mode": "kidney",
                    "products": [],
                    "kidney_rows": [{"cn": "甲", "amount": "1.00"}],
                },
            ),
            (
                "空肾表",
                {"source_mode": "kidney", "products": [], "kidney_rows": []},
            ),
            (
                "空排表",
                {"source_mode": "arrangement", "products": [], "kidney_rows": []},
            ),
        ]
        with tempfile.TemporaryDirectory() as temp_dir:
            for index, (label, spec) in enumerate(cases):
                with self.subTest(label=label):
                    out = Path(temp_dir) / f"invalid-{index}.xlsm"
                    with self.assertRaises(ValueError):
                        build_workbook(spec, out)
                    self.assertFalse(out.exists())

    def test_rejects_non_xlsm_output_extension(self):
        spec = {
            "source_mode": "kidney",
            "products": [],
            "kidney_rows": [{"cn": "甲", "amount": 1}],
        }
        with tempfile.TemporaryDirectory() as temp_dir:
            out = Path(temp_dir) / "kidney.xlsx"
            with self.assertRaisesRegex(ValueError, "必须使用 .xlsm"):
                build_workbook(spec, out)
            self.assertFalse(out.exists())


class FixedPriceAndTemplateTests(unittest.TestCase):
    def test_uses_fixed_price_without_formula_parameters(self):
        spec = {
            "source_mode": "arrangement",
            "title": "固定价",
            "products": [
                {
                    "name": "梨",
                    "allocations": ["A", None, None, None, None, None, None, None, None, None],
                    "price": 12.5,
                }
            ],
            "kidney_rows": [],
        }
        with tempfile.TemporaryDirectory() as temp_dir:
            out = Path(temp_dir) / "fixed-price.xlsm"
            build_workbook(spec, out)
            wb = load_workbook(out, keep_vba=True, data_only=False)
            ws = wb["Sheet1"]

            self.assertEqual(ws["M3"].value, 12.5)
            self.assertEqual(ws["M3"].number_format, "0.00")
            self.assertIsNone(ws["N3"].value)
            self.assertIsNone(ws["O3"].value)
            self.assertEqual(ws["A13"].value, "A")
            self.assertAlmostEqual(ws["B13"].value, 12.5)
            self.assertEqual(ws["C13"].value, 1)

    def test_template_is_sanitized_and_keeps_vba_project(self):
        template = SKILL_ROOT / "assets" / "order-kidney-template.xlsm"
        self.assertTrue(template.is_file())
        with zipfile.ZipFile(template) as archive:
            names = set(archive.namelist())
            self.assertIn("xl/vbaProject.bin", names)
            self.assertFalse(any(name.startswith("xl/media/") for name in names))
            xml_payload = b"\n".join(
                archive.read(name)
                for name in names
                if name.endswith((".xml", ".rels"))
            ).decode("utf-8", errors="ignore")
            for forbidden in (
                "SAMPLE_PERSONAL_NAME",
                "x15ac:absPath",
                "/Users/",
                "支付宝",
            ):
                self.assertNotIn(forbidden, xml_payload)

        wb = load_workbook(template, keep_vba=True, data_only=False)
        ws = wb["Sheet1"]
        self.assertIsNotNone(wb.vba_archive)
        self.assertEqual(len(ws._images), 0)
        self.assertEqual([ws["A12"].value, ws["B12"].value, ws["C12"].value], ["CN", "应付", "点数"])
        self.assertEqual(wb.properties.creator, "goods-skills")
        self.assertEqual(wb.properties.lastModifiedBy, "goods-skills")


class TableFormulaMetadataTests(unittest.TestCase):
    def test_removes_stale_calculated_column_formulas(self):
        spec = {
            "source_mode": "arrangement",
            "title": "公式元数据",
            "products": [
                {
                    "name": "梨",
                    "allocations": ["A", None, None, None, None, None, None, None, None, None],
                    "unit_price": 400,
                    "exchange_rate": 0.0455,
                    "adjustment": 6,
                }
            ],
            "kidney_rows": [],
        }
        with tempfile.TemporaryDirectory() as temp_dir:
            out = Path(temp_dir) / "formula-metadata.xlsm"
            build_workbook(spec, out)
            wb = load_workbook(out, keep_vba=True, data_only=False)
            ws = wb["Sheet1"]
            for table in ws.tables.values():
                for column in table.tableColumns:
                    self.assertIsNone(
                        column.calculatedColumnFormula,
                        f"{table.name}.{column.name} 不应保留旧计算列公式",
                    )


if __name__ == "__main__":
    unittest.main()
